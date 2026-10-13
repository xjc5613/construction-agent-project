# -*- coding: utf-8 -*-
import csv
import json
import mimetypes
from pathlib import Path
from typing import Any, Callable, Dict

from config.settings import MAX_PDF_PAGES, MAX_TABLE_ROWS, MAX_UPLOAD_BYTES

from .models import ParsedDocument, ParsedSection, ParsedTable

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".json", ".pdf", ".docx", ".txt", ".md"}
ALLOWED_MEDIA_TYPES = {
    "text/csv", "text/plain", "text/markdown", "application/json",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/octet-stream",
}


def validate_upload(filename: str, media_type: str, size_bytes: int) -> str:
    safe_name = Path(filename or "").name
    extension = Path(safe_name).suffix.lower()
    if not safe_name or safe_name in {".", ".."} or extension not in ALLOWED_EXTENSIONS:
        raise ValueError("不支持的文件类型")
    if size_bytes <= 0:
        raise ValueError("文件内容为空")
    if size_bytes > MAX_UPLOAD_BYTES:
        raise ValueError(f"文件超过 {MAX_UPLOAD_BYTES} 字节限制")
    guessed = mimetypes.guess_type(safe_name)[0] or "application/octet-stream"
    effective = media_type or guessed
    if effective not in ALLOWED_MEDIA_TYPES and guessed not in ALLOWED_MEDIA_TYPES:
        raise ValueError("文件 MIME 类型不受支持")
    return safe_name


def _parse_csv(path: Path, **_: Any) -> ParsedTable:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 编码无法识别")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    rows = []
    for index, row in enumerate(reader, start=2):
        if len(rows) >= MAX_TABLE_ROWS:
            raise ValueError(f"表格超过 {MAX_TABLE_ROWS} 行限制")
        rows.append({str(key).strip(): value for key, value in row.items() if key is not None} | {"_source_row": index})
    return ParsedTable(columns=list(reader.fieldnames or []), rows=rows)


def _parse_xlsx(path: Path, sheet_name: str = "", **_: Any) -> ParsedTable:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    selected = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    iterator = selected.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
    if not any(headers):
        raise ValueError("工作表缺少表头")
    rows = []
    for source_row, values in enumerate(iterator, start=2):
        if len(rows) >= MAX_TABLE_ROWS:
            raise ValueError(f"表格超过 {MAX_TABLE_ROWS} 行限制")
        if not any(value is not None for value in values):
            continue
        rows.append({
            headers[index]: value for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        } | {"_source_row": source_row})
    return ParsedTable(columns=headers, rows=rows, sheet_name=selected.title)


def _parse_json(path: Path, **_: Any):
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    if isinstance(payload, list):
        if len(payload) > MAX_TABLE_ROWS:
            raise ValueError(f"JSON 数组超过 {MAX_TABLE_ROWS} 行限制")
        rows = []
        for index, item in enumerate(payload, start=1):
            if not isinstance(item, dict):
                raise ValueError("JSON 表格必须是对象数组")
            rows.append(item | {"_source_row": index})
        columns = sorted({key for row in rows for key in row if key != "_source_row"})
        return ParsedTable(columns=columns, rows=rows)
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    return ParsedDocument(
        title=path.stem, text=text,
        sections=[ParsedSection(content=text, section="JSON", ordinal=0)],
    )


def _parse_pdf(path: Path, **_: Any) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"PDF 超过 {MAX_PDF_PAGES} 页限制")
    sections, page_texts = [], []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            page_texts.append(text)
            sections.append(ParsedSection(
                content=text, section=f"第 {index} 页", ordinal=index - 1, page_number=index
            ))
    coverage = len(sections) / max(1, len(reader.pages))
    warnings = []
    if coverage < 0.25 or sum(len(item) for item in page_texts) < 100:
        warnings.append("PDF 文本覆盖率过低，可能需要 OCR")
    title = str((reader.metadata or {}).get("/Title") or path.stem)
    return ParsedDocument(
        title=title, text="\n\n".join(page_texts), sections=sections,
        metadata={"pages": len(reader.pages)}, text_coverage=coverage, warnings=warnings,
    )


def _parse_docx(path: Path, **_: Any) -> ParsedDocument:
    from docx import Document

    document = Document(path)
    sections = []
    ordinal = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            sections.append(ParsedSection(content=text, section="段落", ordinal=ordinal))
            ordinal += 1
    for table_index, table in enumerate(document.tables, start=1):
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if text:
                sections.append(ParsedSection(
                    content=text, section=f"表格 {table_index}", ordinal=ordinal
                ))
                ordinal += 1
    text = "\n".join(section.content for section in sections)
    warnings = [] if text else ["DOCX 未提取到可用文本"]
    return ParsedDocument(
        title=path.stem, text=text, sections=sections,
        text_coverage=1.0 if text else 0, warnings=warnings,
    )


def _parse_text(path: Path, **_: Any) -> ParsedDocument:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("文本编码无法识别")
    paragraphs = [item.strip() for item in text.split("\n\n") if item.strip()]
    sections = [
        ParsedSection(content=item, section="段落", ordinal=index)
        for index, item in enumerate(paragraphs)
    ]
    return ParsedDocument(title=path.stem, text=text.strip(), sections=sections)


PARSERS: Dict[str, Callable[..., Any]] = {
    ".csv": _parse_csv,
    ".xlsx": _parse_xlsx,
    ".json": _parse_json,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".txt": _parse_text,
    ".md": _parse_text,
}


def parse_file(path: Path, **options: Any):
    parser = PARSERS.get(path.suffix.lower())
    if not parser:
        raise ValueError(f"没有适用于 {path.suffix} 的解析器")
    return parser(path, **options)
