from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import json

from openpyxl import load_workbook


path = Path(__file__).resolve().parent / "inputs" / "topic_times.xlsx"
wb = load_workbook(path, read_only=False, data_only=False)
ws = wb["Sheet1"]

style_counts = Counter()
style_rows = defaultdict(list)
fill_counts = Counter()
font_counts = Counter()
row_style_signatures = Counter()
for row in ws.iter_rows():
    sig = []
    for cell in row:
        if cell.value is not None:
            style_counts[cell.style_id] += 1
            style_rows[cell.style_id].append(cell.row)
            fill = cell.fill
            fill_key = (fill.fill_type, fill.fgColor.type, fill.fgColor.rgb, fill.fgColor.indexed)
            fill_counts[fill_key] += 1
            font = cell.font
            font_counts[(font.name, font.sz, font.bold, font.italic, font.color.type if font.color else None, font.color.rgb if font.color and font.color.type == "rgb" else None)] += 1
            sig.append((cell.column, cell.style_id))
    if sig:
        row_style_signatures[tuple(sig)] += 1

check_rows = [1, 2, 100, 300, 500, 700, 900, 1100, 1300, 1500, 1700, 1900, 2100, 2300, 2500, 2700, 2900, 3100, 3300, 3500, 3700, 3900, 4100, 4300, 4500, 4700, 4900, 5100, 5300, 5500, 5700, 5914]
samples = []
for r in check_rows:
    samples.append({
        "row": r,
        "values": [ws.cell(r, c).value for c in range(1, 8)],
        "styles": [ws.cell(r, c).style_id for c in range(1, 8)],
    })

payload = {
    "sheet_state": ws.sheet_state,
    "dimensions": ws.calculate_dimension(),
    "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
    "auto_filter": ws.auto_filter.ref,
    "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
    "hidden_columns": {k: v.hidden for k, v in ws.column_dimensions.items() if v.hidden},
    "hidden_rows_count": sum(1 for v in ws.row_dimensions.values() if v.hidden),
    "defined_names": [n.name for n in wb.defined_names.values()],
    "style_counts": style_counts,
    "style_row_ranges": {str(k): [min(v), max(v), len(set(v))] for k, v in style_rows.items()},
    "fill_counts": {str(k): v for k, v in fill_counts.items()},
    "font_counts": {str(k): v for k, v in font_counts.items()},
    "row_style_signatures_top": [(str(k), v) for k, v in row_style_signatures.most_common(20)],
    "samples": samples,
}

out = Path(__file__).resolve().parent / "extracted" / "workbook_styles.json"
out.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
