from __future__ import annotations

import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
INPUTS = ROOT / "inputs"
OUT = ROOT / "extracted"
OUT.mkdir(parents=True, exist_ok=True)


def extract_docx(path: Path, stem: str) -> dict:
    doc = Document(path)
    paragraphs = []
    for idx, p in enumerate(doc.paragraphs):
        text = re.sub(r"\s+", " ", p.text).strip()
        if text:
            paragraphs.append(
                {
                    "index": idx,
                    "style": p.style.name if p.style else "",
                    "text": text,
                }
            )

    tables = []
    for t_idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([re.sub(r"\s+", " ", c.text).strip() for c in row.cells])
        tables.append({"index": t_idx, "rows": rows})

    images = []
    with ZipFile(path) as zf:
        for name in zf.namelist():
            if name.startswith("word/media/") and not name.endswith("/"):
                target = OUT / f"{stem}_{Path(name).name}"
                target.write_bytes(zf.read(name))
                images.append(str(target.relative_to(ROOT)))

    payload = {
        "source": str(path),
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
        "image_count": len(images),
        "paragraphs": paragraphs,
        "tables": tables,
        "images": images,
    }
    (OUT / f"{stem}.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (OUT / f"{stem}.txt").open("w", encoding="utf-8") as handle:
        for p in paragraphs:
            handle.write(f"[P{p['index']}][{p['style']}] {p['text']}\n")
        for table in tables:
            handle.write(f"\n[TABLE {table['index']}]\n")
            for row in table["rows"]:
                handle.write("\t".join(row) + "\n")
    return payload


def extract_xlsx(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    summary = {"source": str(path), "sheets": []}
    for ws in wb.worksheets:
        sample = []
        nonempty = 0
        min_row = None
        max_row = 0
        min_col = None
        max_col = 0
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            row_values = list(row)
            if any(v is not None for v in row_values):
                nonempty += 1
                min_row = r_idx if min_row is None else min(min_row, r_idx)
                max_row = max(max_row, r_idx)
                occupied = [i + 1 for i, v in enumerate(row_values) if v is not None]
                if occupied:
                    min_col = min(occupied) if min_col is None else min(min_col, min(occupied))
                    max_col = max(max_col, max(occupied))
                if len(sample) < 12:
                    sample.append({"row": r_idx, "values": row_values[:20]})
        summary["sheets"].append(
            {
                "name": ws.title,
                "max_row_reported": ws.max_row,
                "max_col_reported": ws.max_column,
                "nonempty_rows": nonempty,
                "used_bounds": [min_row, min_col, max_row, max_col],
                "sample": sample,
            }
        )
    (OUT / "workbook_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )
    return summary


def main() -> None:
    manuscript = extract_docx(INPUTS / "manuscript.docx", "manuscript")
    ipc = extract_docx(INPUTS / "ipc_network_results.docx", "ipc_network")
    workbook = extract_xlsx(INPUTS / "topic_times.xlsx")
    print(
        json.dumps(
            {
                "manuscript": {
                    "paragraphs": manuscript["paragraph_count"],
                    "tables": manuscript["table_count"],
                    "images": manuscript["image_count"],
                },
                "ipc": {
                    "paragraphs": ipc["paragraph_count"],
                    "tables": ipc["table_count"],
                    "images": ipc["image_count"],
                },
                "workbook": workbook,
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        )
    )


if __name__ == "__main__":
    main()
