from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "output"
FIGURES = ROOT / "figures"
REPORT = OUTPUT / "智能建造关键技术预测成果与发展路线图.docx"

NAVY = "0F2B46"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172033"
MUTED = "5B6573"
LIGHT = "F4F6F9"
LIGHT_BLUE = "E8EEF5"
GRID = "CDD5DF"
WHITE = "FFFFFF"
GOLD = "B87917"


def set_run_font(run, ascii_name="Calibri", east_asia="宋体", size=11, bold=None, color=INK, italic=None):
    run.font.name = ascii_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_fixed_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_para(doc, text="", style=None, align=None, before=0, after=8, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        if style not in ("Title", "Subtitle"):
            set_run_font(r)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p


def add_body(doc, text, lead=None):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.333
    if lead:
        r = p.add_run(lead)
        set_run_font(r, east_asia="微软雅黑", bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=9, color=MUTED)
    return p


def add_figure(doc, path, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    shape = p.add_run().add_picture(str(path), width=Inches(6.35))
    set_alt_text(shape, caption, alt)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_fixed_table_geometry(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(str(text))
        set_run_font(r, east_asia="微软雅黑", size=font_size, bold=True, color=WHITE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        if r_idx % 2 == 1:
            for cell in cells:
                shade_cell(cell, LIGHT)
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.10
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx in (1, len(headers) - 1) else WD_ALIGN_PARAGRAPH.CENTER
            rr = p.add_run(str(value))
            set_run_font(rr, size=font_size, color=INK)
    set_fixed_table_geometry(table, widths)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    remove_paragraph_border(title)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle.paragraph_format.space_after = Pt(20)

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)


def set_page_and_furniture(doc):
    for section_index, section in enumerate(doc.sections):
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = section_index == 0
        header = section.header
        hp = header.paragraphs[0]
        clear_paragraph(hp)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.paragraph_format.space_after = Pt(0)
        r = hp.add_run("智能建造关键技术预测 | 多源证据融合报告")
        set_run_font(r, east_asia="微软雅黑", size=8, color=MUTED)
        footer = section.footer
        fp = footer.paragraphs[0]
        clear_paragraph(fp)
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_after = Pt(0)
        r = fp.add_run("Page ")
        set_run_font(r, size=8, color=MUTED)
        pr = add_field(fp, "PAGE")
        set_run_font(pr, size=8, color=MUTED)
        if section_index == 0:
            clear_paragraph(section.first_page_header.paragraphs[0])
            clear_paragraph(section.first_page_footer.paragraphs[0])


def load_data():
    results = json.loads((OUTPUT / "forecast_results.json").read_text(encoding="utf-8"))
    with (OUTPUT / "theme_annual_counts.csv").open(encoding="utf-8-sig") as handle:
        annual = list(csv.DictReader(handle))
    return results["topics"], results["network_nodes"], annual


def build():
    scores, nodes, annual = load_data()
    doc = Document()
    configure_styles(doc)
    set_page_and_furniture(doc)

    # Cover: editorial_cover with a restrained scientific-report treatment.
    add_para(doc, "TECHNOLOGY FORESIGHT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, before=72, after=18)
    p = doc.add_paragraph(style="Title")
    remove_paragraph_border(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("智能建造领域关键技术预测")
    set_run_font(r, east_asia="微软雅黑", size=30, bold=True, color=NAVY)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于 BERTopic 主题演化与 IPC 动态网络的多源证据融合")
    set_run_font(r, east_asia="微软雅黑", size=15, color=DARK_BLUE)
    add_para(doc, "预测窗口：2027—2040年", align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=6)
    add_para(doc, "证据窗口：2010—2026年（2025完整年度与2026为模型估计）", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "输出：技术优先级、可检验预测、关键IPC候选与技术发展路线图", align=WD_ALIGN_PARAGRAPH.CENTER, after=60)
    p = add_para(doc, "研究用途说明", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    p.runs[0].bold = True
    add_para(doc, "本报告为基于既有论文与本地数据的探索性技术预测，不构成确定性产业承诺。", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "生成日期：2026年10月15日", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_section(WD_SECTION.NEW_PAGE)

    doc.add_heading("摘要", level=1)
    add_body(doc, "本研究将论文中经优化BERTopic得到的九类技术主题，与2015—2024年完整专利活动时序及2010—2025年IPC共现网络结构迁移联合建模。2025年截至10月1日的记录作为右截尾观测，不进入趋势拟合；在此基础上生成2025完整年度与2026年的模型估计，并以80%探索性区间表达不确定性。采用主题动量、近期强度、网络跃迁与聚类融合度四类证据构建综合优先级指数，并在市场牵引、结构跃迁和均衡三种权重情景下进行敏感性检验。")
    add_body(doc, "结果表明，施工全过程数字孪生在三种情景下均保持首位；多专业协同平台、建筑工程大模型以及参数化与生成式设计构成第二梯队。IPC网络显示，G06F30、G06N3、B25J11、G01B11与E02D33等节点具有显著的中心性增长或跨域桥接效应。装配式建筑与建筑规模增材制造的专利基数仍较高，但时序增量趋于平台化，其后续价值更可能来自与机器人、数字孪生、智能材料供应和模块化设计的系统融合。由此提出2027—2040年三阶段路线：2027—2030年完成接口统一与项目级闭环，2031—2035年形成系统闭环并跨项目复制，2036—2040年迈向受控自治、持续合规与循环建造。")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("关键词：")
    set_run_font(r, east_asia="微软雅黑", bold=True, color=NAVY)
    r = p.add_run("智能建造；技术预测；BERTopic；IPC共现网络；数字孪生；建筑工程大模型；技术路线图")
    set_run_font(r)

    doc.add_heading("1 研究数据与预测框架", level=1)
    doc.add_heading("1.1 输入证据与一致性审计", level=2)
    add_body(doc, "预测仅使用用户提供的三项本地输入：论文初稿中的BERTopic主题与生命周期分析、专利主题时间表、以及IPC动态网络结果。论文报告专利样本为5,926件，而当前Excel含5,913条数据行，存在13条差异；该差异可能来自清洗、去重或版本更新，但现有材料不足以确定原因。为避免人为补齐，本报告将其保留为数据一致性限制。")
    add_caption(doc, "表1  输入证据与质量控制")
    add_table(
        doc,
        ["证据ID", "输入", "可用信息", "处理原则"],
        [
            ["E1", "论文初稿", "九个BERTopic主题、模型评价、聚类关系、生命周期与IPC解释", "按原文提取，不改写原始统计值"],
            ["E2", "专利时间表", "5,913条记录；2015—2025阶段标签与申请日", "2015—2024完整观测；2025为截尾观测"],
            ["E3", "IPC网络结果", "阶段1/2/3分别为75、518、979个节点及五类中心性", "统一IPC格式后计算结构跃迁"],
            ["QC", "版本差异", "论文5,926件 vs. Excel 5,913条", "差异13条，列入限制，不推测补齐"],
        ],
        [800, 1700, 3300, 3560],
        font_size=8.2,
    )

    doc.add_heading("1.2 BERTopic主题与时间代理", level=2)
    add_body(doc, "论文中的优化BERTopic模型报告轮廓系数0.682、Calinski–Harabasz指数2896.3；默认参数模型分别为0.415与1245.7。九个主题包括多专业协同平台、参数化与生成式设计、机器人集群协同、施工全过程数字孪生、智能规范审查、建筑工程大模型、装配式建筑、建筑规模增材制造以及模块化与可重构设计。Excel中的既有标签被映射为主题时间代理；其中智能规范审查与模块化/可重构设计缺少专属标签，分别采用可审计的审查/规范词项与模块化/可重构词项作为代理，因此这两类主题的定量置信度较低。")

    doc.add_heading("1.3 多源融合评分", level=2)
    add_body(doc, "综合优先级指数定义为 P = 100 × (0.30M + 0.20I + 0.35N + 0.15C)，其中M为2018—2024年经近期加权的对数动量，I为2021—2024年专利活动强度，N为主题关联IPC节点的中心性水平与阶段跃迁，C为论文层次聚类揭示的融合准备度。各指标经0—1归一化后合成。该指数用于相对排序，不代表技术成功概率。")
    add_body(doc, "为检验权重依赖，另外构造市场牵引情景（更强调动量和专利强度）与结构跃迁情景（更强调IPC网络和主题融合）。报告优先采用均衡情景，并同时给出三种情景的名次区间。证据时间轴以2018—2024年近期加权对数趋势为基础，采用逐年衰减的增长率生成2025完整年度与2026年的模型估计，并输出80%探索性区间。模型值用于延伸证据边界，不是新增观测数据，也不参与原始样本量计算。")

    doc.add_page_break()
    doc.add_heading("2 预测结果", level=1)
    add_figure(
        doc,
        FIGURES / "figure_1_evidence_fusion.png",
        "图1  BERTopic主题时间信号、关键IPC节点迁移与多源证据融合优先级",
        "三联图。A展示2015—2026年九主题年度活动，2025全年与2026采用模型补充；B展示现有三个阶段IPC中心性的迁移，阶段三仍为2021—2025；C展示历史主题动量与网络跃迁，气泡面积正比于2021—2026累计活动量，含2025和2026模型补充。",
    )
    add_body(doc, "图1阅读说明：A的纵轴为ln（年度记录数＋1），用于比较不同规模主题的变化；B的横轴为度中心度，表示节点连接范围，纵轴为中介中心度，表示跨技术桥接作用，数字1、2、3表示原始网络阶段；C的横轴为历史网络跃迁指标，纵轴为历史主题动量，颜色与A图主题一致。C中气泡面积采用线性映射s=0.12×累计活动量，图例给出100、500和1500条的面积参照。气泡表达规模，不表示成功概率；排名与坐标仍采用历史证据，避免将外推值重复作为独立证据。")
    add_caption(doc, "补充表  2021—2026年累计主题活动量（观测与模型补充）")
    add_table(doc, ["主题", "2021—2024观测", "2025全年补充", "2026补充", "2021—2026合计"], [
        [f"{t} {s['topic_name']}", s['recent_count_2021_2024'], s['projection']['2025']['median'], s['projection']['2026']['median'], s['activity_2021_2026']]
        for t, s in scores.items()
    ], [2560, 1700, 1700, 1700, 1700], font_size=7.5)
    add_body(doc, "累计活动量等于2021—2024观测数加2025全年估计及2026估计；2025截至10月1日的部分记录不再重复相加。模型补充属于主题—年度汇总数据，不包含新增专利名称、申请号或IPC共现记录，因此不会改变原始专利样本量，也不能据此重建2026年IPC网络。各主题可重叠，合计不等于去重后的专利总量。证据包同时提供原始输入、年度来源表和补充数据表，供复核。")

    doc.add_heading("2.1 主题优先级", level=2)
    ranking = sorted(scores.items(), key=lambda kv: kv[1]["balanced_score"], reverse=True)
    score_rows = []
    for rank, (topic, s) in enumerate(ranking, start=1):
        score_rows.append([
            rank,
            f"{topic} {s['topic_name']}",
            f"{s['balanced_score']:.1f}",
            f"{s['scenario_rank_range'][0]}–{s['scenario_rank_range'][1]}",
            f"{100*s['annualized_signal_growth']:.1f}%",
            s["recent_count_2021_2024"],
            s["coverage_confidence"],
        ])
    add_caption(doc, "表2  主题综合优先级与敏感性结果")
    add_table(
        doc,
        ["排名", "BERTopic主题", "综合指数", "情景名次", "年化信号动量", "2021—2024记录", "数据置信度"],
        score_rows,
        [560, 2500, 980, 900, 1250, 1250, 1920],
        font_size=7.7,
    )
    add_body(doc, "施工全过程数字孪生（T4）获得94.8分，并在三种情景中均排名第1，表明其优势并非单一权重设定造成。多专业协同平台（T1）、建筑工程大模型（T6）及参数化与生成式设计（T2）在不同情景下位于第2至第4名，构成高潜力融合组。机器人集群协同（T3）具有最大的近期专利活动规模之一，但其增长斜率已由高速扩张转向较稳定的应用深化，因此综合排名处于第5至第6名。")
    add_body(doc, "装配式建筑（T7）与建筑规模增材制造（T8）不是“衰退技术”。其较低排名主要源于新增活动趋稳及部分IPC节点的结构跃迁减弱。例如B33Y70的度中心度由阶段2的0.1257降至阶段3的0.0951，而中介中心度由0.0349升至0.0532，说明其正在从快速扩张的独立主题转向跨技术连接与产业化整合。")

    doc.add_heading("2.2 可检验的活动量预测", level=2)
    proj_rows = []
    for topic in ["T4", "T1", "T6", "T2", "T3", "T7", "T8"]:
        s = scores[topic]
        def fmt(year):
            q = s["projection"][str(year)]
            return f"{q['median']}（{q['low80']}–{q['high80']}）"
        proj_rows.append([f"{topic} {s['topic_name']}", annual[int(topic[1:]) - 1]["2024"], annual[int(topic[1:]) - 1]["2025"], fmt(2025), fmt(2026), s["coverage_confidence"]])
    add_caption(doc, "表3  截至2026年的专利活动证据延伸（中位数及80%区间）")
    add_table(
        doc,
        ["主题", "2024观测", "2025截尾观测", "2025模型估计", "2026模型估计", "解释边界"],
        proj_rows,
        [2200, 950, 1150, 1450, 1450, 2160],
        font_size=7.4,
    )
    add_body(doc, "2024列为完整观测，2025截尾列为截至10月1日的实际记录；2025模型估计与2026模型估计均为生成值，并非新增专利观测。它们是当前检索与主题代理口径下的年度活动信号，而非全球专利族数量。建筑工程大模型和参数化/生成式设计基数较小、区间较宽；智能规范审查和模块化/可重构设计因代理样本不足，不宜给出精确数量预测。因此，本文将数量外推限定为监测基线，核心结论仍是方向、相对优先级及跨主题依赖。")

    doc.add_heading("2.3 关键IPC节点与技术含义", level=2)
    ipc_rows = []
    ipc_specs = [
        ("G06F30", "工程设计三维建模与仿真", "核心跃迁"),
        ("G06N3", "AI/深度学习算法", "核心跃迁"),
        ("B25J11", "建筑机器人操作与机构", "核心跃迁"),
        ("G01B11", "三维激光测量", "使能跃迁"),
        ("G01S17", "激光雷达定位", "使能跃迁"),
        ("B28C7", "智能搅拌与配送", "边缘候选"),
        ("G06Q50", "建造过程数字化管理", "平台核心"),
        ("B66C23", "智能塔吊与吊装", "场景枢纽"),
        ("G16Y10", "建造物联网接口", "标准候选"),
        ("E02D33", "地基监测与预警", "边缘跃迁"),
        ("C08L97", "生物质增强绿色材料", "远期前沿"),
    ]
    for code, meaning, kind in ipc_specs:
        d = nodes.get(code, {})
        ipc_rows.append([
            code,
            meaning,
            f"{d.get('degree', 0):.4f}",
            f"{d.get('betweenness', 0):.4f}",
            f"{d.get('closeness', 0):.4f}",
            kind,
        ])
    add_caption(doc, "表4  阶段3关键IPC节点及预测角色")
    add_table(
        doc,
        ["IPC", "技术含义", "度中心度", "中介中心度", "接近中心度", "预测角色"],
        ipc_rows,
        [1100, 2700, 1250, 1350, 1300, 1660],
        font_size=7.7,
    )

    doc.add_page_break()
    doc.add_heading("3 技术预测结论与机制解释", level=1)
    doc.add_heading("3.1 首要方向：数字孪生将成为施工智能闭环的控制层", level=2)
    add_body(doc, "数字孪生的领先并非仅来自文本热度。G06F30在阶段2至阶段3期间，度中心度由0.0967升至0.1943，中介中心度由0.0533升至0.0745；G01B11的度中心度由0.0077升至0.0746；G06N3由0.0232升至0.1166。三类节点分别对应模型底座、现场感知与智能推理，其同步跃迁支持“感知—建模—预测—控制”链条正在成形。预计2027—2030年主要落点是进度、质量、安全和装备状态的项目级闭环；2031—2035年转向预测性控制与跨项目复用；2036—2040年才可能扩展至项目群—资产连续孪生与实时资源协同。")

    doc.add_heading("3.2 第二梯队：平台、大模型与生成式设计将发生耦合增长", level=2)
    add_body(doc, "T1、T6和T2的共同特征是共享G06F30、G06N3与G06Q50等核心节点。平台层提供可追溯的工程数据和组织协同，大模型层提供知识检索、计划生成与多模态理解，生成式设计层将目标约束转化为可计算方案。短期内，工程大模型更可能以检索增强、规范问答、工艺辅助和风险解释等“受约束助手”形式落地，而不是直接替代责任主体。只有在数据版本、模型来源、规则映射与审计日志形成闭环后，智能审查才可能从辅助核查进入持续合规。")

    doc.add_heading("3.3 机器人从单机替代转向多装备协同", level=2)
    add_body(doc, "B25J11阶段3度中心度达到0.1462、中介中心度达到0.0648，显示机器人技术已进入高连接、高桥接位置；B66C23的度中心度为0.0583、中介中心度为0.0259，体现塔吊和吊装场景的枢纽潜力。未来增长重点不再只是增加机器人品类，而是解决空间定位、任务分配、碰撞规避、设备通信与人机安全边界。预计近期首先在高危、重复和精度敏感工序形成多机协作试点，中期与数字孪生调度融合，远期才可能形成具备自主重规划能力的人机混编施工单元。")

    doc.add_heading("3.4 成熟技术的再定位：装配式与增材制造成为融合载体", level=2)
    add_body(doc, "装配式建筑2021—2024年仍有较高活动规模，但年度记录由267降至145；增材制造同期由91降至53。该变化不应直接解释为技术价值下降，更合理的解释是主题从概念扩散转向工程化筛选和场景整合。未来突破点将集中于三类组合：装配式与机器人装配、增材制造与智能材料配送、以及增材构件与模块化建筑的混合建造。")

    doc.add_heading("3.5 边缘技术机会", level=2)
    add_body(doc, "E02D33的度中心度由阶段2的0.0019升至阶段3的0.0337，且中介中心度升至0.0076，显示地基智能监测与灾害预警具有明确的边缘跃迁特征。G16Y10在阶段3首次进入网络并连接数据处理、通信与设备控制节点，具有接口标准与潜在标准必要专利布局价值。C08L97当前中心性较低，但接近中心度0.2716，适合作为远期绿色材料监测对象。B28C7则可能成为增材制造与现场浇筑之间的材料供应桥梁。")

    doc.add_page_break()
    doc.add_heading("4 技术发展路线图", level=1)
    add_figure(
        doc,
        FIGURES / "figure_2_technology_roadmap.png",
        "图2  智能建造九大主题技术发展路线图（2027—2040）",
        "九条独立技术泳道分别对应T1至T9，分为2027至2030奠基期、2031至2035扩展期、2036至2040转型期。各阶段以浅色卡片呈现；底部独立说明区列出数据支撑、智能与控制、设计与制造三类依赖。A指向B表示A为B提供数据、能力或接口支撑，是规划假设而非已验证因果关系。横带表示贯穿全周期的标准、安全、责任与碳核算约束。",
    )
    add_body(doc, "图2底部的关系符号表示技术支撑方向：T1为T6、T5和T4提供协同数据；T6为T2和T3提供推理能力，T4为T3提供现场状态与控制反馈；T2支撑T9和T8，T9为T7提供模块接口与构型基础。这些关系是结合主题功能与网络关联提出的路线规划假设，不表示已经通过因果检验，也不对应精确发生年份。")
    doc.add_heading("4.1 奠基期：2027—2030年", level=2)
    add_body(doc, "该阶段以接口统一、证据可追溯和项目级闭环为目标。T1建立openBIM/CDE跨角色工作流；T6形成领域RAG、知识图谱与受约束助手；T5推动条款机器可读及规则—模型双路径；T2与T9分别完成约束驱动生成、标准接口和可重构原型；T4形成感知—模型同步及项目级控制；T3开展高危工序多机试点；T7建设柔性制造和构件全链追溯；T8完成材料—工艺—结构联合性能验证。")
    doc.add_heading("4.2 扩展期：2031—2035年", level=2)
    add_body(doc, "该阶段以系统闭环和跨项目复制为目标。T1演化为可信工程数据空间与数字线程；T6支持多模态任务生成、工具调用和可追溯执行；T5进入多模态人机共审；T2将DfMA模块库与审查反馈耦合，T9推进自动构型与快速装拆；T4承担质量、安全和进度的预测性控制；T3实现异构机器人集群调度和孪生驱动重规划；T7打通机器人装配与供应链协同；T8形成搅拌—配送—打印质量闭环。")
    doc.add_heading("4.3 转型期：2036—2040年", level=2)
    add_body(doc, "该阶段以受控自治、持续合规和循环建造为目标。T1支持跨组织知识网络与自治数据治理；T6进入经验证的智能体编排，T5贯通设计—施工—运维持续合规；T2纳入全寿命与循环约束，T9支持建筑全寿命重构和适应性空间；T4扩展为项目群—资产连续孪生；T3形成具备安全边界的人机混编自治单元；T7实现规模化定制、拆解复用与循环建造；T8发展增材—装配混合低碳制造。完全无人化不作为该阶段的默认假设。")

    add_caption(doc, "表5  九大主题的阶段里程碑与依赖关系")
    add_table(
        doc,
        ["主题", "2027—2030", "2031—2035", "2036—2040", "主要前置依赖"],
        [
            ["T1 协同平台", "跨角色工作流", "可信数据空间/数字线程", "跨组织知识网络", "开放标准、数据权属"],
            ["T2 生成式设计", "约束驱动多目标生成", "DfMA与审查反馈闭环", "全寿命循环设计", "T1、T5、T6"],
            ["T3 机器人集群", "多机试点与接口互操作", "异构集群调度", "人机混编受控自治", "T1、T4、T6"],
            ["T4 数字孪生", "项目级控制闭环", "预测性控制", "项目群—资产连续孪生", "T1、感知/测量IPC"],
            ["T5 智能审查", "规则+模型双路径", "多模态人机共审", "全周期持续合规", "T1、T6、法规机器可读"],
            ["T6 工程大模型", "领域RAG与受约束助手", "多模态任务/工具调用", "经验证智能体编排", "T1、模型验证、安全边界"],
            ["T7 装配式建筑", "柔性制造与全链追溯", "机器人装配/供应链协同", "拆解复用与循环建造", "T3、T9"],
            ["T8 增材制造", "材料—工艺—结构验证", "打印质量闭环", "增材×装配低碳制造", "T2、T4、材料标准"],
            ["T9 可重构设计", "标准接口与原型", "自动构型与快速装拆", "全寿命重构", "T2、T7"],
        ],
        [1400, 1850, 1850, 1850, 2410],
        font_size=7.1,
    )

    doc.add_heading("5 监测指标与证伪条件", level=1)
    add_body(doc, "技术预测应被视为可更新假设。若后续数据不能满足以下监测条件，应下调相应路线的优先级或延后时间窗口。")
    add_caption(doc, "表6  路线图的年度监测指标")
    add_table(
        doc,
        ["方向", "领先指标", "支持信号", "证伪/降级条件"],
        [
            ["数字孪生", "G06F30/G01B11/G01S17共现；闭环控制项目数", "连续两年中心性与项目复用上升", "只做可视化、无控制闭环；中心性持续下降"],
            ["工程大模型", "工程数据集、RAG命中率、审查可追溯率", "跨阶段复用且错误可定位", "幻觉率高、无法审计、责任边界不清"],
            ["生成式设计", "约束满足率、方案采纳率、设计周期", "从演示转为工程交付", "采纳率长期低或需大量人工返工"],
            ["机器人集群", "协同任务数、停机率、单位工序成本", "多设备协调优于单机", "接口割裂或安全事件频发"],
            ["智能审查", "规则覆盖率、误报漏报、人工复核量", "规则/模型双路径可解释", "缺少法定认可或漏报不可控"],
            ["增材×装配", "材料闭环控制、构件合格率、混合工法项目", "形成稳定工程标准", "成本/质量不优于传统方案"],
            ["新兴IPC", "节点中心性、跨簇连接、标准提案", "由边缘进入核心或标准网络", "连续两个阶段保持孤立且无场景验证"],
        ],
        [1350, 2850, 2500, 2660],
        font_size=7.7,
    )

    doc.add_heading("6 讨论与局限", level=1)
    add_body(doc, "本研究的主要贡献在于把主题语义、时间动量和网络结构放入同一预测框架，从而避免仅凭高频词或单一中心性判断技术前景。主题聚类揭示技术功能结构，时间序列描述活动变化，IPC网络反映知识组合与桥接位置；三者一致时，预测可信度相对更高。")
    add_body(doc, "局限性主要包括：第一，Excel记录与论文样本数存在13条差异；第二，2025年为部分年度且可能受到专利公开滞后影响，2025完整年度与2026均为模型估计；第三，时间标签来自检索词集合，不等同于BERTopic逐文档主题概率；第四，T5与T9使用词项代理，样本量小；第五，IPC网络中心性反映结构位置，不直接等同于市场规模、技术成熟度或社会收益；第六，外推模型未引入宏观政策、资本投入、成本曲线和专家德尔菲结果；第七，2036—2040路线属于条件性情景，随标准、法规、成本和安全验证进展滚动修订。因此，结果适合用于技术监测与优先级讨论，不宜解释为确定性预测。")

    doc.add_heading("7 结论", level=1)
    add_body(doc, "本研究的多源证据一致支持一个系统性判断：2027—2040年智能建造的主要演化单元将由孤立技术转向可验证、可追溯的技术组合。施工全过程数字孪生在三种权重情景下均位列第一，表明其作为现场状态估计与控制层的相对优先级具有稳健性；多专业协同平台、建筑工程大模型和参数化/生成式设计形成第二梯队，并通过G06F30、G06N3与G06Q50等共享节点产生耦合。机器人集群协同的核心变量将由设备数量转向异构任务协同与安全重规划；装配式建筑、增材制造和可重构设计则更可能作为规模化定制、混合制造与循环建造的工程载体。")
    add_body(doc, "上述结论应按条件性预测解释。2025完整年度与2026的活动量为模型估计，不能替代后续真实观测；2036—2040阶段尤其依赖开放接口、法规机器可读、模型验证、故障降级、全寿命碳核算及责任边界的共同成熟。实践上，研发与政策资源宜优先配置于可信数据底座、项目级闭环验证和跨系统互操作，并通过“年度更新—场景验证—证伪调整”机制，以中心性迁移、跨项目复用、经济性、安全性和合规性持续校正路线图。")

    doc.add_heading("附录A  年度主题代理计数", level=1)
    headers = ["主题"] + [str(y) for y in range(2015, 2025)]
    rows = []
    for item in annual:
        rows.append([f"{item['topic_id']} {item['topic_name']}"] + [item[str(y)] for y in range(2015, 2025)])
    add_table(doc, headers, rows, [2100] + [726] * 10, font_size=7.1)
    add_caption(doc, "注：本表仅列完整观测年度。2025年为截尾观测，2025完整年度与2026的模型估计见表3；T5与T9为词项代理。")

    doc.add_heading("附录B  证据追溯与方法说明", level=1)
    add_body(doc, "E1：论文初稿《基于动态演化的智能建造领域关键技术识别及预测》，用于BERTopic主题、聚类关系、生命周期与IPC解释。E2：专利主题时间表，用于申请年、检索标签与主题代理时序。E3：IPCnetwork结果，用于三个阶段的Degree、Betweenness、Closeness、Eigenvector与Eccentricity指标。所有材料均在本地处理，未上传外部服务。")
    add_body(doc, "方法程序化规范参考：Kassis, T., Agarwal, V., He, Y., Patel, D., & Brueckner, A. M. (2026). Scientific Agent Skills: A Library of Procedural Knowledge for Research Agents. arXiv:2609.00065. https://doi.org/10.48550/arXiv.2609.00065。该文献仅说明科学写作与可视化流程规范，不作为本报告技术预测结论的外部证据。")

    set_page_and_furniture(doc)
    doc.core_properties.title = "智能建造领域关键技术预测成果与发展路线图"
    doc.core_properties.subject = "BERTopic主题演化与IPC动态网络的多源证据融合"
    doc.core_properties.author = "Research analysis workflow"
    doc.core_properties.keywords = "智能建造, BERTopic, IPC网络, 技术预测, 技术路线图"
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    build()
