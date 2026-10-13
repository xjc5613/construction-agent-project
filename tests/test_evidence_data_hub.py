# -*- coding: utf-8 -*-
import importlib
import json
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient

from src.api.project_service import ProjectService
from src.api.forecast_service import create_job, run_forecast_job
from src.api.storage import WorkbenchStorage
from src.ingestion import ImportService, parse_file, validate_upload
from src.library import LibraryDatabase, LibraryRepository
from src.retrieval import EvidenceRetrievalService


@pytest.fixture()
def repository():
    with tempfile.TemporaryDirectory() as directory:
        yield LibraryRepository(LibraryDatabase(Path(directory) / "library.sqlite3"))


def _published(repo, kind, name):
    public_kind = "ipc" if kind.startswith("ipc") else "paper"
    dataset = repo.create_dataset(
        public_kind, name, "1.0",
        technology_node_id="tech_prefabricated_building",
        period_code="2021_2026", public_kind=public_kind,
    )
    if public_kind == "ipc":
        repo.upsert_ipc_metrics(dataset["dataset_id"], [{
            "ipc_code": "E04B", "year": 2021, "patent_count": 1,
        }])
    else:
        repo.add_paper(dataset["dataset_id"], {
            "title": f"{name}基础记录", "year": 2021,
            "keywords": "装配式建筑", "abstract": "基础论文记录",
        })
    return repo.update_dataset_status(dataset["dataset_id"], "published")


def test_database_schema_rollback_and_full_text_search(repository):
    with repository.db.connect() as connection:
        assert connection.execute(
            "SELECT MAX(version) FROM schema_migrations"
        ).fetchone()[0] == 8
        assert "item_id" in {
            row["name"] for row in connection.execute("PRAGMA table_info(evidence_items)")
        }
    with pytest.raises(RuntimeError):
        with repository.db.transaction() as connection:
            connection.execute(
                "INSERT INTO projects(project_id,name) VALUES ('rollback','rollback')"
            )
            raise RuntimeError("rollback")
    assert repository.get_project("rollback") is None

    dataset = _published(repository, "papers", "论文")
    repository.add_paper(dataset["dataset_id"], {
        "title": "装配式建筑智能生产研究", "authors": "研究组", "year": 2024,
        "keywords": "装配式,预制构件", "abstract": "研究预制构件智能生产。",
    }, [{"section": "摘要", "ordinal": 0, "content": "装配式建筑与预制构件生产线。"}])
    results = repository.search_papers("装配式", 10, 2025)
    assert results and results[0]["title"] == "装配式建筑智能生产研究"


def test_table_import_preview_commit_and_trend(repository, tmp_path):
    dataset = repository.create_dataset(
        "ipc", "IPC趋势", "1.0",
        technology_node_id="tech_prefabricated_building",
        period_code="2021_2026", public_kind="ipc",
    )
    service = ImportService(repository)
    content = (
        "IPC分类号,年份,专利数,度中心度\n"
        "E04B,2022,10,0.10\nE04B,2023,15,0.14\nE04B,2024,23,0.19\n"
        "E04B,invalid,30,0.20\n"
    ).encode("utf-8")
    staged = service.stage_file(content, "ipc.csv", "text/csv", "ipc_metrics")
    assert staged["import"]["preview"]["valid_count"] == 3
    assert staged["import"]["preview"]["error_count"] == 1
    committed = service.commit_import(
        staged["import"]["import_id"], dataset["dataset_id"]
    )
    assert committed["committed"] == 3
    repository.update_dataset_status(dataset["dataset_id"], "published")
    rows = repository.ipc_series(["E04B"], 2022, 2024)
    assert [row["year"] for row in rows] == [2022, 2023, 2024]


def test_supported_document_parsers_and_upload_safety(tmp_path):
    text = tmp_path / "paper.md"
    text.write_text("# 标题\n\n装配式建筑研究", encoding="utf-8")
    assert "装配式" in parse_file(text).text

    from docx import Document
    docx_path = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("建筑机器人研究")
    document.save(docx_path)
    assert "建筑机器人" in parse_file(docx_path).text

    from openpyxl import Workbook
    xlsx_path = tmp_path / "ipc.xlsx"
    workbook = Workbook()
    workbook.active.append(["IPC分类号", "年份"])
    workbook.active.append(["E04B", 2024])
    workbook.save(xlsx_path)
    assert parse_file(xlsx_path).rows[0]["IPC分类号"] == "E04B"

    json_path = tmp_path / "papers.json"
    json_path.write_text(
        json.dumps([{"title": "论文", "year": 2024}], ensure_ascii=False),
        encoding="utf-8",
    )
    assert parse_file(json_path).rows[0]["title"] == "论文"

    from pypdf import PdfWriter
    pdf_path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as handle:
        writer.write(handle)
    parsed_pdf = parse_file(pdf_path)
    assert parsed_pdf.text_coverage == 0
    assert parsed_pdf.warnings

    with pytest.raises(ValueError):
        validate_upload("../../payload.exe", "application/octet-stream", 10)


def test_project_isolation_evidence_retrieval_and_snapshot_stability(repository):
    ipc_dataset = _published(repository, "ipc_metrics", "IPC")
    repository.upsert_ipc_metrics(ipc_dataset["dataset_id"], [
        {"ipc_code": "E04B", "year": year, "patent_count": count}
        for year, count in [(2022, 10), (2023, 13), (2024, 18), (2025, 25)]
    ])
    paper_dataset = _published(repository, "papers", "论文")
    for index, year in enumerate((2022, 2023, 2024), start=1):
        repository.add_paper(paper_dataset["dataset_id"], {
            "title": f"装配式建筑研究 {index}", "year": year,
            "keywords": "装配式建筑,预制构件", "abstract": "装配式建筑技术方向研究。",
        }, [{"section": "摘要", "ordinal": 0, "content": "装配式建筑技术与预制构件。"}])

    first = repository.create_project(
        technology_node_id="tech_smart_construction", target_year=2040,
    )
    second = repository.create_project(
        technology_node_id="tech_smart_construction", target_year=2035,
    )
    topic = repository.save_topic({
        "project_id": first["project_id"], "topic_name": "装配式建筑技术",
        "keywords": ["装配式建筑", "E04B"],
        "evidence": [{"ipc": "E04B", "content": "E04B 技术趋势"}],
        "status": "confirmed", "requires_review": False,
        "technologies": [{
            "technology_node_id": "tech_prefabricated_building",
            "is_primary": True, "relevance": 1, "source": "manual",
        }],
    })
    assert repository.list_topics(second["project_id"]) == []

    service = EvidenceRetrievalService(repository)
    package = service.retrieve(first["project_id"], topic["topic_id"], 2022, 2025)
    assert package["quality"]["sufficient"] is False
    assert "历史阶段覆盖" in package["quality"]["missing"]
    original_ids = [item["evidence_id"] for item in package["items"]]
    second_topic = repository.save_topic({
        "project_id": first["project_id"], "topic_name": "装配式建造系统",
        "keywords": ["装配式建筑", "E04B"], "status": "confirmed",
        "requires_review": False,
        "technologies": [{
            "technology_node_id": "tech_prefabricated_building",
            "is_primary": True, "relevance": 1, "source": "manual",
        }],
    })
    second_package = service.retrieve(
        first["project_id"], second_topic["topic_id"], 2022, 2025,
        dataset_ids=[ipc_dataset["dataset_id"], paper_dataset["dataset_id"]],
        min_relevance=0.2,
    )
    assert second_package["query"]["dataset_ids"] == [
        ipc_dataset["dataset_id"], paper_dataset["dataset_id"]
    ]
    assert second_package["query"]["end_year"] == 2025
    assert set(original_ids) & {
        item["evidence_id"] for item in second_package["items"]
    }
    shared_id = next(
        item["evidence_id"] for item in second_package["items"]
        if item["evidence_id"] in original_ids
    )
    service.update_item(shared_id, True, "弱相关", first["project_id"])
    assert service.resolve_evidence(shared_id, first["project_id"])["excluded"] is True
    assert next(
        item for item in repository.get_evidence_package(package["package_id"])["items"]
        if item["evidence_id"] == shared_id
    )["excluded"] is False
    service.update_item(shared_id, False, "", first["project_id"])
    repository.upsert_ipc_metrics(ipc_dataset["dataset_id"], [
        {"ipc_code": "E04B", "year": 2026, "patent_count": 35}
    ])
    unchanged = repository.get_evidence_package(package["package_id"])
    assert [item["evidence_id"] for item in unchanged["items"]] == original_ids


def test_project_api_is_scoped_and_rejects_unsafe_file(
    repository, monkeypatch, tmp_path
):
    app_module = importlib.import_module("src.api.app")
    library_routes = importlib.import_module("src.api.library_routes")
    project_routes = importlib.import_module("src.api.project_routes")
    monkeypatch.setattr(library_routes, "repository", repository)
    monkeypatch.setattr(library_routes, "import_service", ImportService(repository))
    monkeypatch.setattr(project_routes, "repository", repository)
    monkeypatch.setattr(project_routes, "project_service", ProjectService(repository))
    monkeypatch.setattr(
        project_routes, "retrieval_service", EvidenceRetrievalService(repository)
    )
    monkeypatch.setattr(app_module, "storage", WorkbenchStorage(tmp_path / "workbench"))
    client = TestClient(app_module.app)
    project = client.post("/api/projects", json={
        "technology_node_id": "tech_smart_construction",
        "target_year": 2040,
        "name": "API项目",
    }).json()
    status_response = client.patch(
        f"/api/projects/{project['project_id']}/status",
        json={"status": "material_ready"},
    )
    assert status_response.status_code == 200
    assert status_response.json()["status"] == "material_ready"
    response = client.post(
        f"/api/projects/{project['project_id']}/materials/text",
        json={"title": "材料", "content": "E04B 装配式建筑技术研究"},
    )
    assert response.status_code == 201
    unsafe = client.post(
        f"/api/projects/{project['project_id']}/materials/upload",
        files={"file": ("payload.exe", b"unsafe", "application/octet-stream")},
    )
    assert unsafe.status_code == 422
    topic = client.post(
        f"/api/projects/{project['project_id']}/topics",
        json={"topic_name": "装配式建筑"},
    ).json()
    archived = client.post(
        f"/api/projects/{project['project_id']}/topics/{topic['topic_id']}/archive"
    )
    assert archived.json()["status"] == "archived"
    restored = client.post(
        f"/api/projects/{project['project_id']}/topics/{topic['topic_id']}/restore"
    )
    assert restored.json()["status"] == "candidate"
    legacy = client.post(
        "/api/documents",
        json={"title": "兼容材料", "content": "兼容接口材料", "source_type": "text"},
    )
    assert legacy.status_code == 201
    assert legacy.headers["Deprecation"] == "true"


def test_material_can_discover_non_seed_topic(repository):
    project = repository.create_project("新主题项目")
    service = ProjectService(repository)
    material = service.add_text_material(
        project["project_id"],
        "自适应施工声学监测技术研究",
        "自适应施工声学监测技术通过声纹分析识别设备异常，并形成现场预警方法。",
    )
    topics = service.extract_topics(project["project_id"], [material["material_id"]])
    assert topics
    assert topics[0]["topic_name"] != "待校准技术主题"


def test_evidence_snapshot_to_three_node_result(repository, tmp_path):
    project = repository.create_project("路线图项目")
    evidence_id = "evidence_ipc_demo"
    package = {
        "package_id": "package_demo",
        "project_id": project["project_id"],
        "topic_id": "topic_demo",
        "quality": {"score": 90, "sufficient": True, "missing": []},
        "items": [{
            "evidence_id": evidence_id, "evidence_type": "ipc_metric",
            "source_id": "ipc_demo", "title": "E04B · 2025",
            "content": "E04B 度中心度持续增长", "excluded": False,
            "metadata": {"ipc_code": "E04B", "year": 2025, "degree_centrality": 0.18},
        }],
    }
    topic = {
        "topic_id": "topic_demo", "topic_name": "装配式建筑技术",
        "keywords": ["装配式", "E04B"], "selected": True,
        "ipc_evidence": [{"ipc": "E04B", "trend": "增长", "quote": "增长", "document_id": evidence_id}],
        "fusion_pairs": [], "sample_abstracts": ["装配式建筑论文方向持续发展"],
        "evidence_package": package,
    }
    storage = WorkbenchStorage(tmp_path / "jobs")
    job = create_job(
        storage, [topic["topic_id"]],
        {"enable_per_topic_roadmap": True, "enable_reasoning_chain": True},
        project_id=project["project_id"], topic_snapshot=[topic],
        evidence_snapshot={"project_id": project["project_id"], "packages": [package]},
    )
    round1 = [{
        "topic_name": topic["topic_name"], "bottlenecks_2030_2035": ["标准协同"],
        "breakthroughs_by_2040": ["全生命周期闭环"], "evidence_ids": [evidence_id],
        "ipc_inference": ["IPC持续增长"], "paper_inference": ["研究方向活跃"],
        "evidence_synthesis": ["双通道收敛"], "evidence_conflicts": [],
    }]
    round3 = [
        {"year": 2030, "category": "技术里程碑", "description": "标准验证", "evidence_ids": [evidence_id]},
        {"year": 2035, "category": "产品里程碑", "description": "规模应用", "evidence_ids": [evidence_id]},
        {"year": 2040, "category": "技术里程碑", "description": "成熟闭环", "evidence_ids": [evidence_id]},
    ]
    roadmaps = [{
        "topic_name": topic["topic_name"],
        "roadmap": {
            year: {"stage_description": f"{year}阶段", "milestones": [{
                "name": f"{year}里程碑", "description": "证据约束结果",
                "trl_level": trl, "uncertainty_level": uncertainty,
                "evidence_ids": [evidence_id],
            }]}
            for year, trl, uncertainty in [
                ("2025", 3, "low"), ("2030", 5, "medium"),
                ("2035", 7, "medium"), ("2040", 9, "high"),
            ]
        },
    }]
    with patch("src.api.forecast_service.run_round1", return_value=round1), \
         patch("src.api.forecast_service.run_round2", return_value={"new_paradigm_name": "证据协同建造"}), \
         patch("src.api.forecast_service.run_round3", return_value=round3), \
         patch("src.api.forecast_service.run_all_topics_roadmap", return_value=roadmaps), \
         patch("src.api.forecast_service.generate_job_assets", return_value={"assets": [], "failed_assets": []}):
        run_forecast_job(storage, job["job_id"])
    saved = storage.get_job(job["job_id"])
    result = storage.get_job_artifact(job["job_id"], "result.json")
    charts = storage.get_job_artifact(job["job_id"], "charts.json")
    assert saved["status"] == "completed"
    assert [item["year"] for item in result["round3_result"]] == [2030, 2035, 2040]
    assert result["evidence_snapshot_checksum"]
    assert charts["ipc_trends"][0]["evidence_id"] == evidence_id
    assert (storage.job_dir(job["job_id"]) / "evidence_snapshot.json").exists()
